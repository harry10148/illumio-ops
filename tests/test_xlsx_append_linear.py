"""append_df_rows must not look rows up through max_row/max_column.

Both properties scan every cell already in the sheet, so calling them once
per appended row made XLSX export O(rows²): a 10k-flow report spent about
eight minutes in this one helper (measured 2026-10). Asserting the property
is never read is deterministic, unlike a timing threshold.
"""
import pandas as pd
import pytest
from openpyxl import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from src.report.exporters.xlsx_exporter import append_df_rows


def test_append_never_scans_the_sheet(monkeypatch):
    def _boom(self):
        raise AssertionError("max_row/max_column scans the whole sheet")

    df = pd.DataFrame({"a": [1, 2, 3], "b": ["x", "blocked", "y"]})
    wb = Workbook()
    ws = wb.active
    ws.append(["preamble"])  # rows already there: numbering must continue after them
    monkeypatch.setattr(Worksheet, "max_row", property(_boom))
    monkeypatch.setattr(Worksheet, "max_column", property(_boom))
    append_df_rows(ws, df)
    monkeypatch.undo()

    assert [c.value for c in ws[2]] == ["a", "b"]
    assert ws.cell(row=2, column=1).font.b
    assert ws.cell(row=3, column=1).value == 1
    # the alert row ("blocked") is filled across the table's columns
    assert ws.cell(row=4, column=1).fill.fgColor.rgb == ws.cell(row=4, column=2).fill.fgColor.rgb
    assert ws.cell(row=4, column=1).fill.fill_type == "solid"
    assert ws.cell(row=3, column=1).fill.fill_type in (None, "none")
    assert ws.max_row == 5
