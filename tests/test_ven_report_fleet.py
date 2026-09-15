"""Task 5 — VEN 狀態報表的四個車隊章節。

Spec: docs/superpowers/specs/2026-09-09-ven-fleet-manager-design.md
Plan: docs/superpowers/plans/2026-09-09-ven-fleet-manager.md Task 5

報表與 GUI 讀的是**同一支** `analyze_fleet`，所以兩邊的數字不可能各說各話；
這裡守的是「那份分析有沒有真的走到紙上」，不是分析本身（那在
tests/test_fleet_analysis.py）。
"""
import types
from unittest.mock import MagicMock

import pytest


def _wl(host, mode="selective", *, version="26.2.20-2063", status="active",
        hb=0.2, labels=None, os_id="ubuntu-x86_64-22.04"):
    return {
        "href": "/orgs/1/workloads/" + host,
        "hostname": host, "os_id": os_id, "enforcement_mode": mode,
        "interfaces": [{"address": "10.0.0.1"}],
        "labels": labels if labels is not None else [{"key": "app", "value": "web"},
                                                     {"key": "env", "value": "prod"}],
        "agent": {"status": {
            "status": status, "hours_since_last_heartbeat": hb,
            "security_policy_sync_state": "active",
            "last_heartbeat_on": "2026-09-14T00:00:00Z",
            "security_policy_received_at": "2026-09-14T00:00:00Z",
            "agent_version": version, "agent_health": [],
            "agent_health_errors": {"errors": [], "warnings": []},
        }},
    }


WORKLOADS = [
    _wl("idle-1", mode="idle"),
    _wl("vis-1", mode="visibility_only"),
    _wl("sel-1"),
    _wl("sel-2", version="23.4.11-8"),
    _wl("off-1", hb=99.0, labels=[]),
]


def _cm(target=None):
    settings = {"timezone": "UTC"}
    if target is not None:
        settings["fleet_target_ven_version"] = target
    return types.SimpleNamespace(config={"settings": settings})


def _generate(tmp_path, *, target=None, lang="en", workloads=WORKLOADS):
    from src.report.ven_status_generator import VenStatusGenerator
    api = MagicMock()
    api.fetch_managed_workloads.return_value = workloads
    gen = VenStatusGenerator(_cm(target), api_client=api)
    return gen.generate(output_dir=str(tmp_path), lang=lang)


# ── 分析有沒有進到 module_results ──────────────────────────────────────────

def test_the_report_carries_the_same_fleet_analysis_the_gui_reads(tmp_path):
    r = _generate(tmp_path)
    fleet = r.module_results.get("fleet")
    assert fleet, "報表沒有 fleet 分析"
    assert fleet["total"] == len(WORKLOADS)
    assert set(fleet["pipeline"]) and "health_score" in fleet


def test_the_target_version_comes_from_settings(tmp_path):
    r = _generate(tmp_path, target="26.2.20-2063")
    v = r.module_results["fleet"]["versions"]
    assert v["target"] == "26.2.20-2063"
    assert v["on_target"] == 4        # 五台裡四台在目標版本
    assert v["needs_upgrade"] == 1


def test_no_target_means_no_upgrade_number_in_the_report(tmp_path):
    """報表不得憑空生出一個「待升級 0 台」——沒設目標就沒有這個數字。"""
    r = _generate(tmp_path)
    v = r.module_results["fleet"]["versions"]
    assert v["target"] is None
    assert v["needs_upgrade"] is None


def test_the_fleet_analysis_does_not_cost_a_second_fetch(tmp_path):
    from src.report.ven_status_generator import VenStatusGenerator
    api = MagicMock()
    api.fetch_managed_workloads.return_value = WORKLOADS
    VenStatusGenerator(_cm(), api_client=api).generate(output_dir=str(tmp_path))
    assert api.fetch_managed_workloads.call_count == 1


# ── 四個章節有沒有走到 HTML ────────────────────────────────────────────────

SECTION_IDS = ("fleet-pipeline", "fleet-compat", "fleet-score", "fleet-gaps")


def _html(tmp_path, **kw):
    """實際產出報表 HTML。

    走 export() 而不是私有的 _build()：要驗的是「讀者打開檔案會看到什麼」，
    而 export() 才是產品真正走的那條路。
    """
    import pathlib
    from src.report.exporters.ven_html_exporter import VenHtmlExporter
    r = _generate(tmp_path, **kw)
    lang = kw.get("lang", "en")
    out = VenHtmlExporter(r.module_results, df=r.dataframe, lang=lang).export(
        output_dir=str(tmp_path))
    return pathlib.Path(out).read_text(encoding="utf-8")


@pytest.mark.parametrize("lang", ["en", "zh_TW"])
def test_all_four_chapters_reach_the_document(tmp_path, lang):
    doc = _html(tmp_path, lang=lang)
    missing = [i for i in SECTION_IDS if ('id="%s"' % i) not in doc]
    assert missing == [], "這些章節沒有出現在報表裡：%s" % missing


@pytest.mark.parametrize("lang", ["en", "zh_TW"])
def test_every_chapter_is_reachable_from_the_contents(tmp_path, lang):
    """章節存在但目錄沒列＝讀者找不到。這個 repo 有過一模一樣的漏。"""
    doc = _html(tmp_path, lang=lang)
    missing = [i for i in SECTION_IDS if ('#%s"' % i) not in doc]
    assert missing == [], "目錄沒有這幾章：%s" % missing


def test_a_partial_score_says_so_on_the_page(tmp_path):
    """沒設目標版本時分數是部分的——紙上必須看得出來，不能只印一個數字。"""
    from src.i18n import t
    doc = _html(tmp_path)
    assert t("rpt_ven_fleet_partial", lang="en") in doc


# ── xlsx ─────────────────────────────────────────────────────────────────────

def test_the_workbook_carries_every_workload_not_a_sample(tmp_path):
    """HTML 的樣本表截到 50 列；xlsx 是拿去篩選排序的，截斷比沒有還糟。"""
    from openpyxl import load_workbook
    from src.i18n import t
    from src.report.ven_status_generator import generate_ven_xlsx

    # 超過 HTML 的 50 列樣本上限，否則這支測不出「有沒有截斷」——注入實測：
    # 只有 5 台的 fixture 下，把 xlsx 也截到 50 依然全綠。
    many = WORKLOADS + [_wl("bulk-%03d" % i) for i in range(70)]
    r = _generate(tmp_path, workloads=many)
    assert len(r.module_results["fleet"]["workloads_index"]) > 50
    out = str(tmp_path / "ven.xlsx")
    generate_ven_xlsx(r.module_results, out, lang="en")
    wb = load_workbook(out)
    sheet_name = t("rpt_xlsx_sheet_ven_fleet", lang="en")
    assert sheet_name in wb.sheetnames, wb.sheetnames
    ws = wb[sheet_name]
    # 標頭一列 + 每台一列
    assert ws.max_row == len(r.module_results["fleet"]["workloads_index"]) + 1


def test_an_old_snapshot_without_fleet_still_produces_a_workbook(tmp_path):
    """舊快照沒有 fleet 鍵——不得因此炸掉整份 xlsx。"""
    from openpyxl import load_workbook
    from src.report.ven_status_generator import generate_ven_xlsx

    r = _generate(tmp_path)
    results = dict(r.module_results)
    results.pop("fleet", None)
    out = str(tmp_path / "old.xlsx")
    generate_ven_xlsx(results, out, lang="en")
    assert load_workbook(out).sheetnames
