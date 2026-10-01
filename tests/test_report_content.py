"""報表內容：第一頁、佐證、enforcement 進度、XLSX／CSV、圖表無障礙。"""
from __future__ import annotations

import zipfile
from types import SimpleNamespace

import pandas as pd
from openpyxl import load_workbook

from src.report.analysis import mod03_uncovered_flows, mod12_executive_summary
from src.report.analysis.mod_enforcement_progress import analyze as enforcement_analyze
from src.report.exporters._exec_summary import render_exec_summary_html
from src.report.exporters.html_exporter import _format_evidence


def _finding(sev, rid="L010"):
    return SimpleNamespace(rule_id=rid, rule_name="Cross-Env Lateral", severity=sev,
                           description="SMB between prod and dev", recommendation="Block SMB across envs",
                           category="LateralMovement", evidence={})


def _summary(findings):
    df = pd.DataFrame([{"src_ip": "1", "dst_ip": "2", "src_app": "a", "dst_app": "b", "src_managed": True,
                        "dst_managed": True, "port": 445, "proto": "TCP", "num_connections": 1,
                        "policy_decision": "potentially_blocked"}] * 10)
    results = {"mod01": {"total_flows": 10, "src_managed_pct": 100},
               "mod03": mod03_uncovered_flows.uncovered_flows(df),
               "mod04": {}, "mod08": {}, "mod11": {}, "mod15": {}, "findings": findings}
    return mod12_executive_summary.executive_summary(results)


def test_first_screen_puts_risk_kpis_and_critical_action_first():
    out = _summary([_finding("CRITICAL")])
    keys = [k["label_key"] for k in out["kpis"][:4]]
    assert keys[:2] == ["mod12_kpi_maturity_score", "mod12_kpi_crit_high_findings"]
    assert "mod12_kpi_total_flows" not in keys
    assert out["top_actions"][0]["severity"] == "CRITICAL"
    assert "1 critical" in out["verdict"]
    html = render_exec_summary_html(out, report_name="x", include_heading=False)
    assert "Block SMB across envs" in html and html.index("exec-actions") < html.index("kpi-strip")


def test_evidence_is_not_cut_and_unlabeled_is_named():
    html = _format_evidence({"top_env_pairs": str({("prod", "dev"): 168, ("", "backup"): 14}),
                             "items": str([f"item-number-{i}-with-a-long-name-xxxxxxxxxxxxxxxxxxxx" for i in range(7)])},
                            lang="en")
    assert "prod → dev: 168" in html
    assert "(unlabeled) → backup" in html
    assert "item-number-0-with-a-long-name-xxxxxxxxxxxxxxxxxxxx" in html
    assert "(+2 more)" in html


def test_breaks_on_enforcement_suggests_allow_rule():
    df = pd.DataFrame([
        dict(src_ip="10.0.0.1", src_app="web", src_env="prod", src_enforcement="full",
             dst_ip="10.0.1.1", dst_app="db", dst_env="prod", dst_enforcement="visibility_only",
             port=5432, proto="TCP", num_connections=5, policy_decision="potentially_blocked"),
        dict(src_ip="10.0.0.9", src_app="web", src_env="prod", src_enforcement="full",
             dst_ip="10.0.1.1", dst_app="db", dst_env="prod", dst_enforcement="visibility_only",
             port=443, proto="TCP", num_connections=9, policy_decision="allowed"),
    ])
    out = enforcement_analyze(df)
    rules = out["breaks_on_enforcement"]["Suggested Allow Rule"].tolist()
    assert rules == ["web (prod) → db (prod) : 5432/TCP"]
    progress = out["progress"].set_index("App (Env)")
    assert progress.loc["db (prod)", "Enforced %"] == 0.0
    assert progress.loc["web (prod)", "Enforced %"] == 100.0


def test_xlsx_has_numbers_findings_raw_and_no_pb_red_fill(tmp_path):
    from src.report.report_generator import build_traffic_xlsx
    mr = {"mod12": {"kpis": [{"label": "Total Flows", "value": "1,234"},
                             {"label": "Coverage", "value": "45.2%"}]},
          "mod02": {"summary": pd.DataFrame([{"Decision": "potentially_blocked", "Flows": 3},
                                             {"Decision": "blocked", "Flows": 1}])},
          "findings": [_finding("HIGH")]}
    raw = pd.DataFrame([{"src_ip": "10.0.0.1", "num_connections": 3, "extra": {"k": "v"}}])
    path = build_traffic_xlsx(mr, str(tmp_path / "r.xlsx"), profile="security_risk", raw_df=raw)
    wb = load_workbook(path)
    kpi = wb[wb.sheetnames[1]]
    assert kpi["B2"].value == 1234 and kpi["B3"].value == 45.2
    assert "Findings" in wb.sheetnames and "Raw Flows" in wb.sheetnames
    assert wb["Raw Flows"].auto_filter.ref
    pd_ws = [ws for ws in wb.worksheets if ws.title.startswith("Policy")][0]
    fills = {r[0].value: r[0].fill.fgColor.rgb for r in pd_ws.iter_rows(min_row=3, max_row=4)}
    assert fills["potentially_blocked"] != fills["blocked"]


def test_csv_zip_has_findings_and_manifest(tmp_path):
    from src.report.exporters.csv_exporter import CsvExporter
    path = CsvExporter({"findings": [_finding("HIGH")],
                        "mod01": {"top_ports": pd.DataFrame([{"Port": 443, "Flows": 3}])}}).export(str(tmp_path))
    names = zipfile.ZipFile(path).namelist()
    assert "findings.csv" in names and "_manifest.csv" in names


def test_chart_svg_has_accessible_name():
    from src.report.exporters.chart_renderer import render_matplotlib_svg
    svg = render_matplotlib_svg({"type": "bar", "title": "Ports",
                                 "data": {"labels": ["a"], "values": [1]}})
    assert svg.startswith('<svg role="img" aria-label="Ports"') and "<title>Ports</title>" in svg
