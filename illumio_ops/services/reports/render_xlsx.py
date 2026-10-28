"""ReportDocument -> XLSX: a summary sheet (key metrics, findings, data quality) and one sheet per table with
every row (up to XLSX_ROW_CAP, stated when hit). Text that a spreadsheet would run as a formula is prefixed
with an apostrophe."""
from __future__ import annotations

from datetime import tzinfo
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

from illumio_ops.domain.report import XLSX_ROW_CAP, Cell, ReportDocument, period_text
from illumio_ops.domain.values import Language
from illumio_ops.services.reports.render_html import fmt_time, kpi_trend, tr
from illumio_ops.services.reports.sheets import RESERVED, sheet_names

_FORMULA_START = ("=", "+", "-", "@", "\t", "\r")
MAX_WIDTH = 60
WRAP = Alignment(wrap_text=True, vertical="top")
WRAP_FROM = MAX_WIDTH - 10  # characters; shorter text fits its column
BOLD = Font(bold=True)
SUMMARY_SHEET = {Language.EN: RESERVED[0], Language.ZH_TW: RESERVED[1]}


def safe(value: Cell) -> Cell:
    if isinstance(value, str) and value.startswith(_FORMULA_START):
        return "'" + value
    return value


def _width(value: object) -> int:
    return min(len(str(value)) if value is not None else 0, MAX_WIDTH)


def _sheet(wb: Workbook, title: str, rows: list[list[Cell]], *, bold_rows: set[int], freeze_after: int | None) -> None:
    """Write-only sheets keep memory flat on large tables; widths are measured before writing."""
    ws = wb.create_sheet(title)
    widths: dict[int, int] = {}
    for row in rows:
        for i, v in enumerate(row):
            widths[i] = max(widths.get(i, 8), _width(v))
    for i, w in widths.items():
        ws.column_dimensions[get_column_letter(i + 1)].width = min(w + 2, MAX_WIDTH)
    if freeze_after is not None:
        ws.freeze_panes = f"A{freeze_after + 1}"
    # A styled cell costs ~50x a plain value in openpyxl: only text that can exceed its column is wrapped (a
    # column is at most MAX_WIDTH wide), so 100,000-row tables stay seconds, not minutes.
    for n, row in enumerate(rows, start=1):
        out: list[Any] = []
        for v in row:
            value = safe(v)
            if n in bold_rows:
                cell = WriteOnlyCell(ws, value=value)
                cell.font = BOLD
                out.append(cell)
            elif isinstance(v, str) and len(v) > WRAP_FROM:
                cell = WriteOnlyCell(ws, value=value)
                cell.alignment = WRAP
                out.append(cell)
            else:
                out.append(value)
        ws.append(out)


def _summary_rows(doc: ReportDocument, tz: tzinfo | None) -> list[list[Cell]]:
    lang = doc.lang
    rows: list[list[Cell]] = [
        [doc.title], [tr(lang, "rpt_cover_pce"), doc.pce_label],
        [tr(lang, "rpt_cover_period"), period_text(doc.period_start, doc.period_end, tz)],
        [tr(lang, "rpt_cover_generated"), fmt_time(doc.generated_at.isoformat(), tz)], [],
        [tr(lang, "rpt_xlsx_metric"), tr(lang, "rpt_xlsx_value"), tr(lang, "rpt_xlsx_previous"),
         tr(lang, "rpt_xlsx_change"), tr(lang, "rpt_xlsx_note")],
    ]
    rows += [[k.label, k.value if k.value is not None else tr(lang, "rpt_kpi_no_value"),
              k.previous, kpi_trend(k, lang), k.note] for k in doc.kpis]
    rows += [[], [tr(lang, "rpt_findings_title")]]
    findings: list[list[Cell]] = [[tr(lang, "severity_" + f.severity.value), f.title, f.why, f.action]
                                  for f in doc.findings]
    rows += findings or [[tr(lang, "rpt_findings_none")]]
    rows += [[], [tr(lang, "rpt_data_quality_title")]]
    rows += [[q] for q in doc.data_quality]
    return rows


def render_xlsx(doc: ReportDocument, *, tz: tzinfo | None) -> bytes:
    wb = Workbook(write_only=True)
    _sheet(wb, SUMMARY_SHEET[doc.lang], _summary_rows(doc, tz), bold_rows={1, 6}, freeze_after=None)
    tables = doc.tables()
    for (_, t), name in zip(tables, sheet_names([t.title for _, t in tables]), strict=True):
        # Excel caps a sheet name at 31 characters: a shortened one has the table's whole title above it.
        notes: list[list[Cell]] = [[t.title]] if name != t.title else []
        if t.limit_note:
            notes.append([t.limit_note])
        if len(t.rows) > XLSX_ROW_CAP:
            notes.append([tr(doc.lang, "rpt_xlsx_capped", shown=XLSX_ROW_CAP, total=len(t.rows))])
        header: list[Cell] = [c.label for c in t.columns]
        body: list[list[Cell]] = [
            [fmt_time(v, tz) if col.kind == "time" and isinstance(v, str) else v
             for col, v in zip(t.columns, row, strict=True)]
            for row in t.rows[:XLSX_ROW_CAP]] or [[t.empty_text]]
        _sheet(wb, name, [*notes, header, *body], bold_rows={len(notes) + 1}, freeze_after=len(notes) + 1)
    buf = BytesIO()
    wb.save(buf)
    return buf.getvalue()
