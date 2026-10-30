"""Every built report type in both languages (P5a Task 10): no raw message key in the HTML, the XLSX or the email,
and a zh_TW report carries no English sentence of the report's own copy."""
import json
import re
from datetime import UTC, datetime, timedelta
from io import BytesIO
from pathlib import Path

import pytest
from openpyxl import load_workbook

from illumio_ops.domain.filter import EndpointFilter, FlowCondition, LabelCond
from illumio_ops.domain.findings_rules import RULES
from illumio_ops.domain.values import DataOrigin, Language, ReportType, RhcState
from illumio_ops.pce.policy import PendingItem, PolicyNames
from illumio_ops.services.analysis.policy_resolver import Resolved
from illumio_ops.services.analysis.predicted_block import (
    PredictedGroup,
    PredictedSummary,
)
from illumio_ops.services.analysis.ransomware import HostRansomware
from illumio_ops.services.reports import policy_diff as policy_diff_report
from illumio_ops.services.reports import policy_resolver as policy_resolver_report
from illumio_ops.services.reports import ransomware_report, rule_hit_count
from illumio_ops.services.reports import readiness as readiness_report
from illumio_ops.services.reports.registry import BuildInput, ReportParams, specs
from illumio_ops.services.reports.render_email import render_email
from illumio_ops.services.reports.render_html import render_html
from illumio_ops.services.reports.render_xlsx import render_xlsx
from illumio_ops.services.rhc import RhcPull, RhcRow, RhcStatus
from illumio_ops.store.flowdb import analysis
from illumio_ops.store.flowdb.db import db_time
from illumio_ops.store.flowdb.events import upsert_events
from illumio_ops.store.flowdb.flows import upsert_flows
from tests_v6.services.collect.conftest import configured_ctx
from tests_v6.services.reports.test_audit import event as audit_event
from tests_v6.services.reports.test_findings import flow
from tests_v6.services.reports.test_lateral_posture import (
    _detection,
    _exposure,
    _hosts,
    _snapshot,
)
from tests_v6.services.reports.test_new_connections import _baseline
from tests_v6.services.reports.test_unmanaged import _rollup

NOW = datetime(2026, 10, 27, 1, 0, tzinfo=UTC)
START, END = datetime(2026, 10, 20, tzinfo=UTC), datetime(2026, 10, 27, tzinfo=UTC)
RAW_KEY = re.compile(r"\b(?:rpt|gui_reports|detector_name|decision|severity)_[a-z0-9_]+\b")
# PCE console words zh_TW keeps in English (spec §6.2.1): not report copy left untranslated.
PCE_TERMS = frozenset({"Rule Hit Count", "Override Deny"})
APP_DB = FlowCondition(dst=EndpointFilter(labels=(LabelCond(key="app", values=("DB",)),)))
EN = json.loads((Path(__file__).resolve().parents[3] / "illumio_ops/i18n/en.json")
                .read_text(encoding="utf-8"))


@pytest.fixture(scope="module")
def ctx(tmp_path_factory):
    c = configured_ctx(tmp_path_factory.mktemp("i18n"))
    c.settings.update(c.settings.current().revision,
                      lambda d: d["lateral"].update(critical_assets=[{"key": "app", "value": "Finance"}]))
    _hosts(c)
    _exposure(c)
    _snapshot(c, "2026-10-26")
    _detection(c, "fanout", datetime(2026, 10, 21, tzinfo=UTC), severity="high")
    _detection(c, "auth_to_non_dc", datetime(2026, 10, 22, tzinfo=UTC), summary={"target": "dc", "port": 88})
    d = datetime(2026, 10, 22, tzinfo=UTC)
    _baseline(c, [("/orgs/1/workloads/a", "/orgs/1/workloads/db", 3389, d, "allowed"),
                  ("ip:10.9.9.9", "/orgs/1/workloads/a", 445, d, "unknown")])
    _rollup(c, [("2026-10-21", "ip:10.5.5.5", "/orgs/1/workloads/a", 445, "allowed", 10),
                ("2026-10-23", "/orgs/1/workloads/a", "ip:8.8.8.8", 53, "allowed", 3)])
    _rule_flows(c)
    _p5b2_sample(c)
    mp = pytest.MonkeyPatch()
    # The readiness report's draft recompute is a PCE call: a fixed answer here, so every section has content.
    mp.setattr(readiness_report, "predicted_summary", lambda ctx, spec, **kw: PredictedSummary(
        groups=[PredictedGroup(dst_app="DB", rule_href="/orgs/1/sec_policy/draft/rule_sets/1/deny_rules/2", flows=3,
                               connections=7)], rules={}, blocked_flows=3, blocked_connections=7, undetermined_flows=2, capped=True,
        computed_at=NOW, unsupported=False, missing_rules=["/orgs/1/sec_policy/draft/rule_sets/1/deny_rules/2"]))
    _rhc_sample(mp)
    _pdf_sample(mp)
    mp.setattr(policy_resolver_report, "resolve", lambda ctx, **kw: Resolved(
        rows=[{"ruleset_name": "Web", "rule_href": "/r/1", "action": "allow", "src_ip": "<5001 hosts: expansion capped>",
               "dst_ip": "10.0.0.2", "port": 1433, "protocol": "TCP", "truncated": "src"}],
        per_ruleset={"Web": 1}, rule_count=1, rules_capped=1, truncated_rulesets=["Web"], rows_omitted=3,
        not_expanded=[("Web", "/r/2", "virtual_services"), ("Web", "/r/3", "unknown_reference")],
        computed_at=NOW))
    mp.setattr(ransomware_report, "host_ransomware", lambda ctx, href, **kw: HostRansomware.model_construct(
        breakdown=None, allow_rules={"445/6": []}, allow_status={"445/6": "error"}))
    yield c
    mp.undo()


def _rhc_sample(mp):
    """Rule Hit Count on, one rule hit, one at 0, one with no record, one pending change (P5c-2)."""
    a = "/orgs/1/sec_policy/active/rule_sets/1"
    mp.setattr(rule_hit_count, "rhc_status", lambda ctx: RhcStatus(
        state=RhcState.ENABLED, pce_report=True, ven_scopes=True, reason_key=None, detail=None, write_key=True,
        saas=False))
    mp.setattr(rule_hit_count, "pull_rhc", lambda ctx, start, end: RhcPull(rows=[
        RhcRow(rule_href=f"{a}/sec_rules/1", ruleset_name="Web", description="", hit_count=9,
               last_hit_at="2026-10-26T10:00:00Z"),
        RhcRow(rule_href=f"{a}/sec_rules/2", ruleset_name="Web", description="", hit_count=0, last_hit_at=None),
        RhcRow(rule_href=f"{a}/sec_rules/9", ruleset_name="Old", description="", hit_count=1, last_hit_at=None)],
        start=start, end=end, unparsed=1))
    rule = {"consumers": [], "providers": [], "ingress_services": [], "enabled": True}
    mp.setattr(rule_hit_count, "fetch_rule_sets", lambda client, pversion, *, deadline: [
        {"href": a, "name": "Web", "rules": [{**rule, "href": f"{a}/sec_rules/{i}"} for i in (1, 2, 3)],
         "deny_rules": [{**rule, "href": f"{a}/deny_rules/4", "override": True}]}])
    mp.setattr(rule_hit_count, "fetch_policy_names",
               lambda client, pversion, *, deadline: PolicyNames(labels={}, names={}))
    mp.setattr(rule_hit_count, "fetch_pending", lambda client, *, deadline: [PendingItem(
        kind="rule_sets", href="/orgs/1/sec_policy/draft/rule_sets/7", name="New", update_type="create",
        updated_at="2026-10-26T00:00:00Z")])


def _pdf_sample(mp):
    """Every change kind, field and risk the policy diff report words (P5c-3)."""
    from illumio_ops.domain.policy_diff import Change
    from illumio_ops.services.analysis.policy_diff import Attribution, PolicyDiff
    d = "/orgs/1/sec_policy/draft"
    rule_fields = ("rule_type", "enabled", "description", "providers", "consumers", "services")
    obj_fields = {"ip_list": ("ip_ranges", "fqdns", "description"),
                  "service": ("service_ports", "windows_services", "description"),
                  "label_group": ("labels", "sub_groups", "description")}
    changes = [Change(kind="rule_set", change="modified", ruleset="Q", name="Q", ref="1", href=f"{d}/rule_sets/1",
                      field=f, draft="False" if f == "enabled" else "x", active="True" if f == "enabled" else "y",
                      risk="high" if f == "enabled" else "") for f in ("name", "enabled", "description", "rule_count")]
    changes += [Change(kind="rule_set", change=ch, ruleset="N", name="N", ref="2", href=f"{d}/rule_sets/2", field="*",
                       draft="N" if ch == "added" else "", active="N" if ch == "removed" else "", rule_count=3,
                       risk="medium" if ch == "removed" else "") for ch in ("added", "removed")]
    changes += [Change(kind="rule", change="modified", ruleset="Q", name="#5", ref="deny_rules/5",
                       href=f"{d}/rule_sets/1/deny_rules/5", field=f,
                       draft="override_deny" if f == "rule_type" else "a", active="deny" if f == "rule_type" else "b",
                       rule_type="override_deny") for f in rule_fields]
    changes += [Change(kind=k, change="modified", ruleset="", name="O", ref="3", href=f"{d}/{k}s/3", field=f,
                       draft="a, b", active="a", scope_expanded=True, risk="high") for k, fs in obj_fields.items()
                for f in fs]
    mp.setattr(policy_diff_report, "policy_diff", lambda ctx, **kw: PolicyDiff(
        changes=changes, pending=[], attribution={f"{d}/rule_sets/1": Attribution(
            actor="admin", at=datetime(2026, 10, 25, tzinfo=UTC), event_type="rule_set.update")},
        other_pending=[PendingItem(kind="firewall_settings", href=f"{d}/firewall_settings", name=None,
                                   update_type="update", updated_at="2026-10-25T00:00:00Z")],
        attribution_limited=True, computed_at=NOW))


def _p5b2_sample(c):
    """VEN status, events, ransomware and tampering, so the P5b-2 reports render every section (P5b-2 Task 6)."""
    hb = NOW - timedelta(hours=5)
    with c.flows().connect() as conn:
        conn.execute("UPDATE workloads SET agent_status = 'active', last_heartbeat_at = ?, policy_sync_state = "
                     "'syncing', fw_config_current = 0, policy_received_at = ?, policy_applied_at = ?",
                     (db_time(NOW - timedelta(minutes=5)), db_time(NOW - timedelta(minutes=10)), db_time(hb)))
        conn.execute("UPDATE workloads SET online = 0, last_heartbeat_at = ? WHERE href = '/orgs/1/workloads/b'",
                     (db_time(hb),))
        analysis.write_result(conn, "tamper", params={}, at=NOW, result={"unmatched": 0, "hosts": [
            {"href": "/orgs/1/workloads/a", "hostname": "app-01", "n_60d": 4, "n_7d": 2, "pattern": "periodic",
             "median_minutes": 10.0, "last_at": NOW.isoformat(), "unexplained_7d": 1, "has_signal": True}]})
        svc = {"name": "SMB", "port": 445, "proto": 6, "severity": "critical", "category": "admin",
               "protection_state": "unprotected", "port_status": "listening", "active_policy": "allowed", "weight": 40,
               "coverage": 0.0, "contribution": 40.0, "traffic": "priority", "recent_sources": 3}
        listed = {"host": "/orgs/1/workloads/a", "hostname": "app-01", "name": "SMB", "port": 445, "proto": 6,
                  "severity": "critical", "weight": 40, "recent_sources": 3}
        analysis.write_result(conn, "ransomware", params={}, at=NOW, result={
            "unsupported": False, "contradictions": 1, "priority": [listed], "safe_to_block": [listed],
            "suspicious": [listed], "hosts": [
                {"href": "/orgs/1/workloads/a", "hostname": "app-01", "enforcement_mode": "selective",
                 "pce_percent": 10.0, "pce_exposure": "critical", "computed_percent": 10.0, "approximate": True,
                 "contradiction": False, "services": [svc]},
                {"href": "/orgs/1/workloads/b", "hostname": "bak-01", "enforcement_mode": "selective",
                 "pce_percent": 0.0, "pce_exposure": "high", "computed_percent": 100.0, "approximate": True,
                 "contradiction": True, "services": [{**svc, "protection_state": "protected_open"}]}]})
        conn.execute("INSERT INTO ransomware_snapshot_daily (day, workload_href, name, port, proto, severity, "
                     "protection_state, port_status) VALUES ('2026-10-22', '/orgs/1/workloads/a', 'SMB', 445, 6, "
                     "'critical', 'unprotected', 'listening')")
        conn.execute("INSERT INTO org_risk_snapshot_daily (day, summary_json) VALUES ('2026-10-26', ?)",
                     (json.dumps({"ransomware": {"ransomware_protection_coverage_percent": 12.5}}),))
        d = datetime(2026, 10, 22, 3, tzinfo=UTC)
        events = [audit_event("agent.tampering", d, system=True),
                  audit_event("sec_policy.create", d, user="alice", extra={"resource_changes": [{"resource": {
                      "sec_policy": {"commit_message": "open db", "workloads_affected": 3}}}]}),
                  audit_event("rule_set.update", d, user="alice"),
                  *(audit_event("request.authentication_failed", d + timedelta(minutes=i), user="eve",
                                ip="10.1.1.1", status="failure") for i in range(5))]
        upsert_events(conn, events, origin=DataOrigin.INGEST, now=NOW)
        conn.commit()


def _rule_flows(c):
    """Flows that make every findings rule fire (thresholds lowered), so all rule copy is rendered."""
    c.settings.update(c.settings.current().revision, lambda d: d["reports"]["rule_thresholds"].update(
        unmanaged_connection=1, min_policy_coverage_pct=100, lateral_outbound_dst=1, user_destination=1,
        cross_env_connection=1, discovery_protocol=1, db_unique_src_app=1, identity_unique_src_app=1, blast_radius=1,
        unmanaged_critical=1, pb_lateral=1, exfil_mb=1, cross_env_lateral=1))
    flows = [flow(port=445, src_env="Dev", dst_env="Prod"), flow(port=5900), flow(port=22, decision="potentially_blocked"),
             flow(port=1433, unmanaged_src=True), flow(port=443, unmanaged_src=True),
             flow(src="jump", dst="t1", port=22), flow(src="jump", dst="t2", port=22),
             flow(dst="t1", port=443, user="CORP\\alice"), flow(dst="t2", port=443, user="CORP\\alice"),
             flow(port=23, decision="blocked"), flow(port=137, proto=17), flow(port=1433, src_app="A", src_env="Prod"), flow(port=1433),
             flow(port=389, src_app="A"), flow(port=389, src_app="B"), flow(port=443, unmanaged_dst=True, bytes_=2 * 1024 * 1024),
             *(flow(dst=f"p{i}", port=443, bytes_=1000) for i in range(20)),
             flow(dst="big", port=443, bytes_=50 * 1024 * 1024)]
    with c.flows().connect() as conn:
        upsert_flows(conn, flows, origin=DataOrigin.INGEST, now=NOW)


def _custom(ctx):
    """A custom report with each kind of section (P6e): a flow ranking, a flow detail, an event ranking."""
    from illumio_ops.domain.segment import GroupBy
    from illumio_ops.services import custom_reports
    from illumio_ops.services.workbench import EventRequest, FlowRequest
    window = {"start": END - timedelta(days=1), "end": END}
    rep = custom_reports.add_section(ctx, report_id=None, name="Custom", title="Ranking", by="t", now=NOW,
                                     request=FlowRequest(mode="flows", view="ranking",
                                                         group_by=(GroupBy(dimension="dst_host"),), **window))
    custom_reports.add_section(ctx, report_id=rep.id, name=None, title="Detail", by="t", now=NOW,
                               request=FlowRequest(mode="flows", view="detail", **window))
    custom_reports.add_section(ctx, report_id=rep.id, name=None, title="Events", by="t", now=NOW,
                               request=EventRequest(mode="events", view="ranking",
                                                    group_by=(GroupBy(dimension="event_type"),), **window))
    return rep.id


def _outputs(ctx, report_type, lang):
    cond = APP_DB if specs()[report_type].requires_app else None
    custom_id = _custom(ctx) if report_type is ReportType.CUSTOM else None
    inp = BuildInput(ctx=ctx, params=ReportParams(report_type=report_type, filter=cond, custom_id=custom_id),
                     start=START, end=END, lang=lang, tz=UTC, now=NOW, pce_label="pce")
    doc = specs()[report_type].build(inp).model_copy(update={"gui_link": "https://ops.example/reports?id=1"})
    wb = load_workbook(BytesIO(render_xlsx(doc, tz=UTC)))
    cells = "\n".join(str(c.value) for ws in wb.worksheets for row in ws.iter_rows() for c in row
                      if c.value is not None)
    mail = render_email(doc, tz=UTC)
    return {"html": render_html(doc, tz=UTC), "xlsx": cells + "\n" + "\n".join(wb.sheetnames),
            "email": mail.subject + "\n" + mail.text + "\n" + (mail.html or "")}


@pytest.mark.parametrize("lang", [Language.EN, Language.ZH_TW])
@pytest.mark.parametrize("report_type", list(ReportType))
def test_no_raw_keys_in_output(ctx, report_type, lang):
    for kind, text in _outputs(ctx, report_type, lang).items():
        body = text.split("<style>", 1)[0] + text.split("</style>", 1)[-1].split("<script>", 1)[0]
        found = sorted(set(RAW_KEY.findall(body)))
        assert found == [], f"{report_type} {lang} {kind}: {found}"


@pytest.mark.parametrize("report_type", list(ReportType))
def test_zh_tw_has_no_english_report_copy(ctx, report_type):
    english = [v for k, v in EN.items() if k.startswith(("rpt_", "detector_name_")) and len(v) > 12
               and not re.search(r"\{\w+\}", v) and v.isascii() and v not in PCE_TERMS]
    for kind, text in _outputs(ctx, report_type, Language.ZH_TW).items():
        leaked = [s for s in english if s in text]
        assert leaked == [], f"{report_type} {kind}: {leaked[:3]}"


def test_period_kept_in_both_languages(ctx):
    for lang in Language:
        assert "2026-10-20 – 2026-10-26" in _outputs(ctx, ReportType.LATERAL_POSTURE, lang)["email"]
    assert START + timedelta(days=7) == END


def test_every_rule_rendered(ctx):
    """The sample makes all rules fire, so the two checks above read every rule's copy in both languages."""
    for lang in Language:
        doc = specs()[ReportType.SECURITY_RISK].build(BuildInput(
            ctx=ctx, params=ReportParams(report_type=ReportType.SECURITY_RISK), start=START, end=END, lang=lang,
            tz=UTC, now=NOW, pce_label="pce"))
        rows = next(s for s in doc.sections if s.id == "findings").blocks[0].rows
        assert {r[0] for r in rows} == {r.rule_id for r in RULES}


def test_injected_missing_key_is_caught(ctx, monkeypatch):
    from illumio_ops import i18n
    real = i18n.dictionary

    def without(lang):
        return {k: v for k, v in real(lang).items() if k != "rpt_rule_b002_name"}
    monkeypatch.setattr(i18n, "dictionary", without)
    html = _outputs(ctx, ReportType.SECURITY_RISK, Language.ZH_TW)["html"]
    assert "rpt_rule_b002_name" in RAW_KEY.findall(html.split("</style>", 1)[-1])
