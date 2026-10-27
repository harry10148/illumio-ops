from datetime import UTC
from io import BytesIO

from openpyxl import load_workbook

from illumio_ops.domain.report import Section
from illumio_ops.services.reports.render_xlsx import render_xlsx
from tests_v6.services.reports.conftest import make_doc, table


def _wb(doc):
    return load_workbook(BytesIO(render_xlsx(doc, tz=UTC)))


def test_xlsx_one_sheet_per_table_all_rows(doc_with_rows):
    wb = _wb(doc_with_rows(450))
    data = [ws for ws in wb.worksheets if ws.title != "Summary"]
    assert len(data) == 1 and data[0].max_row == 451  # header + every row


def test_xlsx_summary_sheet_has_kpis_and_quality():
    ws = _wb(make_doc())["Summary"]
    values = [c.value for row in ws.iter_rows() for c in row if c.value is not None]
    assert "Pivot hosts" in values and 3 in values and "Traffic covers the whole period." in values


def test_xlsx_neutralises_formulas(doc_with_rows):
    wb = _wb(doc_with_rows(1, cell="=cmd|' /C calc'!A0"))
    ws = [w for w in wb.worksheets if w.title != "Summary"][0]
    assert ws["A2"].value.startswith("'=")


def test_xlsx_sheet_names_unique_and_short():
    long = "A very long table title that exceeds thirty-one characters: [x]"
    doc = make_doc(sections=(Section(id="a", title="A", blocks=(table([("a", 1)], title=long),
                                                                table([("b", 2)], title=long))),))
    names = _wb(doc).sheetnames
    assert len(set(names)) == len(names) and all(len(n) <= 31 for n in names)
    assert not any(ch in n for n in names for ch in "[]:*?/\\")


def test_xlsx_builder_limit_note_written():
    t = table([("a", 1)], total=5).model_copy(update={"limit_note": "Only the first 1 of 5"})
    ws = [w for w in _wb(make_doc(sections=(Section(id="s", title="S", blocks=(t,)),))).worksheets
          if w.title != "Summary"][0]
    assert ws["A1"].value == "Only the first 1 of 5" and ws["A2"].value == "Col 0"


def test_xlsx_long_cell_kept_whole(doc_with_rows):
    ws = [w for w in _wb(doc_with_rows(1, cell="y" * 2500)).worksheets if w.title != "Summary"][0]
    assert ws["A2"].value == "y" * 2500 and ws["A2"].alignment.wrap_text
